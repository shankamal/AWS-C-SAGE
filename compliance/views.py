import csv
import io
import json
from datetime import datetime, timezone

from django.contrib import messages
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.shortcuts import redirect, render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .aws.orchestrator import ComplianceOrchestrator
from .aws.suppressions import SuppressionStore
from .models import Finding, ResourceInventory, ScanRun
from .storage import store


ORANGE = "F37021"
NAVY = "003366"
DARK_NAVY = "052F5F"
LIGHT_HEADER = "EAF0F6"
LIGHT_ORANGE = "FFF1E8"
WHITE = "FFFFFF"
BORDER = "D8E0E8"


def healthz(request):
    return JsonResponse({"status": "ok", "service": "C-SAGE", "storage": "flat-file"})


def _latest():
    return ScanRun.objects.first()


def dashboard(request):
    scan = _latest()
    findings = Finding.objects.for_scan(scan)
    modules = []
    for name in sorted({f.module for f in findings}):
        module_findings = [f for f in findings if f.module == name]
        modules.append({
            "name": name,
            "total": len(module_findings),
            "non_compliant": sum(1 for f in module_findings if f.status == Finding.Status.NON_COMPLIANT),
            "suppressed": sum(1 for f in module_findings if f.status == Finding.Status.SUPPRESSED),
        })
    recent = [f for f in findings if f.status != Finding.Status.COMPLIANT][:12]
    return render(request, "compliance/dashboard.html", {"scan": scan, "recent_findings": recent, "modules": modules})


def findings_view(request):
    scan = _latest()
    base = Finding.objects.for_scan(scan)
    account = request.GET.get("account", "")
    module = request.GET.get("module", "")
    status = request.GET.get("status", "")
    keyword = request.GET.get("q", "").strip().lower()

    findings = base
    if account:
        findings = [f for f in findings if f.account_id == account]
    if module:
        findings = [f for f in findings if f.module == module]
    if status:
        findings = [f for f in findings if f.status == status]
    if keyword:
        findings = [
            f for f in findings
            if keyword in f.resource_id.lower()
            or keyword in f.title.lower()
            or keyword in f.rule_id.lower()
            or keyword in f.details.lower()
        ]

    accounts = sorted({(f.account_id, f.account_name) for f in base}, key=lambda item: item[1])
    module_names = sorted({f.module for f in base})
    ctx = {
        "scan": scan,
        "findings": findings[:1000],
        "accounts": accounts,
        "module_names": module_names,
        "filters": {"account": account, "module": module, "status": status, "q": request.GET.get("q", "")},
    }
    return render(request, "compliance/findings.html", ctx)


def _inventory_filters(request, rows):
    account = request.GET.get("account", "")
    service = request.GET.get("service", "")
    resource_type = request.GET.get("resource_type", "")
    region = request.GET.get("region", "")
    status = request.GET.get("status", "")
    keyword = request.GET.get("q", "").strip().lower()

    result = rows
    if account:
        result = [r for r in result if r.account_id == account]
    if service:
        result = [r for r in result if r.service == service]
    if resource_type:
        result = [r for r in result if r.resource_type == resource_type]
    if region:
        result = [r for r in result if r.region == region]
    if status:
        result = [r for r in result if r.status == status]
    if keyword:
        result = [r for r in result if any(
            keyword in str(value).lower()
            for value in (r.resource_name, r.resource_id, r.resource_arn, r.service, r.resource_type, r.account_name, r.account_id)
        )]
    return result, {
        "account": account,
        "service": service,
        "resource_type": resource_type,
        "region": region,
        "status": status,
        "q": request.GET.get("q", ""),
    }


def inventory_view(request):
    scan = _latest()
    base = ResourceInventory.objects.for_scan(scan)
    filtered, filters = _inventory_filters(request, base)
    filtered.sort(key=lambda r: (r.account_name.lower(), r.service.lower(), r.resource_type.lower(), r.resource_name.lower(), r.resource_id.lower()))

    coverage = [row for row in store.read_inventory_coverage() if scan and row.get("scan_run_id") == scan.id]
    coverage_summary = {}
    for row in coverage:
        key = (row.get("service", "Unknown"), row.get("resource_type", "Unknown"))
        item = coverage_summary.setdefault(key, {
            "service": key[0], "resource_type": key[1], "resource_count": 0,
            "complete": 0, "access_denied": 0, "errors": 0, "regions": set(), "accounts": set(),
        })
        item["resource_count"] += int(row.get("resource_count", 0) or 0)
        status = row.get("scan_status", "")
        if status == "COMPLETE":
            item["complete"] += 1
        elif status == "ACCESS_DENIED":
            item["access_denied"] += 1
        else:
            item["errors"] += 1
        item["regions"].add(row.get("region", ""))
        item["accounts"].add(row.get("account_id", ""))
    coverage_rows = []
    for item in coverage_summary.values():
        item["regions"] = len(item["regions"])
        item["accounts"] = len(item["accounts"])
        coverage_rows.append(item)
    coverage_rows.sort(key=lambda x: (x["service"].lower(), x["resource_type"].lower()))

    paginator = Paginator(filtered, 200)
    page = paginator.get_page(request.GET.get("page", 1))
    accounts = sorted({(r.account_id, r.account_name) for r in base}, key=lambda x: x[1].lower())
    services = sorted({r.service for r in base} | {r.get("service", "") for r in coverage if r.get("service")})
    resource_types = sorted({r.resource_type for r in base} | {r.get("resource_type", "") for r in coverage if r.get("resource_type")})
    regions = sorted({r.region for r in base} | {r.get("region", "") for r in coverage if r.get("region")})
    statuses = sorted({r.status for r in base if r.status})

    query_without_page = request.GET.copy()
    query_without_page.pop("page", None)
    ctx = {
        "scan": scan,
        "page": page,
        "filtered_count": len(filtered),
        "total_count": len(base),
        "coverage": coverage_rows,
        "coverage_count": len(coverage_rows),
        "zero_resource_count": sum(1 for x in coverage_rows if x["resource_count"] == 0),
        "coverage_issue_count": sum(1 for x in coverage_rows if x["access_denied"] or x["errors"]),
        "accounts": accounts,
        "services": services,
        "resource_types": resource_types,
        "regions": regions,
        "statuses": statuses,
        "filters": filters,
        "query_without_page": query_without_page.urlencode(),
        "export_query": request.GET.urlencode(),
    }
    return render(request, "compliance/inventory.html", ctx)


def resource_detail(request, resource_key):
    scan = _latest()
    resource = ResourceInventory.objects.by_key(scan, resource_key)
    if not resource:
        messages.error(request, "Resource not found in the latest Master Inventory snapshot.")
        return redirect("inventory")

    pretty = {
        "tags": json.dumps(resource.tags, indent=2, sort_keys=True, default=str),
        "configuration": json.dumps(resource.configuration, indent=2, sort_keys=True, default=str),
        "networking": json.dumps(resource.networking, indent=2, sort_keys=True, default=str),
        "security": json.dumps(resource.security, indent=2, sort_keys=True, default=str),
        "relationships": json.dumps(resource.relationships, indent=2, sort_keys=True, default=str),
        "raw_attributes": json.dumps(resource.raw_attributes, indent=2, sort_keys=True, default=str),
    }
    return render(request, "compliance/resource_detail.html", {"scan": scan, "resource": resource, "pretty": pretty})


@require_POST
def run_scan(request):
    try:
        actor = request.user.get_username() if getattr(request, "user", None) else "system"
        scan = ComplianceOrchestrator().run(actor or "system")
        messages.success(request, f"Audit scan #{scan.id} completed. Master Inventory contains {scan.resource_count:,} resources.")
    except Exception as exc:
        messages.error(request, f"Audit scan failed: {exc}")
    return redirect("dashboard")


@require_POST
def toggle_suppression(request, finding_id):
    scan = _latest()
    findings = Finding.objects.for_scan(scan)
    finding = next((item for item in findings if item.id == finding_id), None)
    if not finding:
        messages.error(request, "Finding not found in the latest flat-file scan results.")
        return redirect("findings")

    suppressions = SuppressionStore()
    try:
        actor = request.user.get_username() if getattr(request, "user", None) else "system"
        if finding.status == Finding.Status.SUPPRESSED:
            suppressions.unsuppress(finding.finding_key, actor)
            finding.status = Finding.Status.NON_COMPLIANT
            messages.success(request, "Finding unsuppressed.")
        else:
            reason = request.POST.get("reason", "Suppressed through C-SAGE UI")[:500]
            suppressions.suppress(finding.finding_key, actor, reason, finding=finding)
            finding.status = Finding.Status.SUPPRESSED
            messages.success(request, "Finding suppressed and persisted with resource details in the local flat-file store.")

        store.write_findings([item.to_dict() for item in findings])
        if scan:
            scan.compliant_count = sum(1 for item in findings if item.status == Finding.Status.COMPLIANT)
            scan.non_compliant_count = sum(1 for item in findings if item.status == Finding.Status.NON_COMPLIANT)
            scan.suppressed_count = sum(1 for item in findings if item.status == Finding.Status.SUPPRESSED)
            scan.save()
    except Exception as exc:
        messages.error(request, f"Suppression update failed: {exc}")
    return redirect(request.META.get("HTTP_REFERER", "/findings/"))


def _compact_json(value):
    return json.dumps(value or {}, separators=(",", ":"), sort_keys=True, default=str)


def _flatten(value, prefix=""):
    if isinstance(value, dict):
        if not value:
            yield prefix or "$", "{}"
        for key, child in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            yield from _flatten(child, path)
    elif isinstance(value, list):
        if not value:
            yield prefix or "$", "[]"
        for index, child in enumerate(value):
            path = f"{prefix}[{index}]" if prefix else f"[{index}]"
            yield from _flatten(child, path)
    else:
        text = "" if value is None else str(value)
        if len(text) <= 32000:
            yield prefix or "$", text
        else:
            for index in range(0, len(text), 32000):
                yield f"{prefix or '$'} [part {index // 32000 + 1}]", text[index:index + 32000]


def _style_header(ws, row=1):
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(color=WHITE, bold=True)
    border = Border(bottom=Side(style="thin", color=BORDER))
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[row].height = 24


def _format_inventory_sheet(ws):
    _style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "A": 18, "B": 16, "C": 26, "D": 16, "E": 22, "F": 28, "G": 28, "H": 44,
        "I": 18, "J": 22, "K": 24, "L": 48, "M": 48, "N": 48, "O": 48, "P": 48, "Q": 20,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _filtered_inventory_for_export(request):
    scan = _latest()
    rows = ResourceInventory.objects.for_scan(scan)
    rows, _ = _inventory_filters(request, rows)
    return scan, rows


def export_inventory_csv(request):
    scan, rows = _filtered_inventory_for_export(request)
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="c-sage-master-inventory.csv"'
    writer = csv.writer(response)
    writer.writerow([
        "Resource Key", "Account ID", "Account Name", "Region", "Service", "Resource Type", "Resource Name",
        "Resource ID", "ARN", "Status", "Creation Time", "Tags", "Networking", "Security", "Attached/Dependent Resources",
        "Configuration Details", "Discovery Source"
    ])
    for r in rows:
        writer.writerow([
            r.resource_key, r.account_id, r.account_name, r.region, r.service, r.resource_type, r.resource_name,
            r.resource_id, r.resource_arn, r.status, r.creation_time, _compact_json(r.tags), _compact_json(r.networking),
            _compact_json(r.security), _compact_json(r.relationships), _compact_json(r.configuration), r.discovery_source,
        ])
    return response


def export_inventory_xlsx(request):
    scan, rows = _filtered_inventory_for_export(request)
    coverage = [row for row in store.read_inventory_coverage() if scan and row.get("scan_run_id") == scan.id]

    wb = Workbook()
    summary = wb.active
    summary.title = "Inventory Summary"
    summary["A1"] = "C-SAGE Complete AWS Master Inventory"
    summary["A1"].font = Font(size=18, bold=True, color=WHITE)
    summary["A1"].fill = PatternFill("solid", fgColor=ORANGE)
    summary.merge_cells("A1:D1")
    summary["A3"] = "Scan ID"
    summary["B3"] = scan.id if scan else ""
    summary["A4"] = "Generated At (UTC)"
    summary["B4"] = datetime.now(timezone.utc).isoformat()
    summary["A5"] = "Resources Exported"
    summary["B5"] = len(rows)
    summary["A6"] = "Coverage Records"
    summary["B6"] = len(coverage)
    summary["A7"] = "Zero-resource Coverage Records"
    summary["B7"] = sum(1 for x in coverage if int(x.get("resource_count", 0) or 0) == 0)
    summary["A8"] = "Access/Error Coverage Records"
    summary["B8"] = sum(1 for x in coverage if x.get("scan_status") != "COMPLETE")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 40
    for cell in summary["A"][2:8]:
        cell.font = Font(bold=True, color=NAVY)

    ws = wb.create_sheet("Master Inventory")
    ws.append([
        "Resource Key", "Account ID", "Account Name", "Region", "Service", "Resource Type", "Resource Name",
        "Resource ID", "ARN", "Status", "Creation Time", "Tags (JSON)", "Networking (JSON)", "Security (JSON)",
        "Attached/Dependent Resources (JSON)", "Configuration Details (JSON)", "Discovery Source"
    ])
    for r in rows:
        ws.append([
            r.resource_key, r.account_id, r.account_name, r.region, r.service, r.resource_type, r.resource_name,
            r.resource_id, r.resource_arn, r.status, str(r.creation_time or ""), _compact_json(r.tags),
            _compact_json(r.networking), _compact_json(r.security), _compact_json(r.relationships),
            _compact_json(r.configuration), r.discovery_source,
        ])
    _format_inventory_sheet(ws)

    cov = wb.create_sheet("Service Coverage")
    cov.append(["Account ID", "Account Name", "Region", "Service", "Resource Type", "Resource Count", "Scan Status", "Message", "Source"])
    for row in sorted(coverage, key=lambda x: (x.get("account_name", ""), x.get("service", ""), x.get("resource_type", ""), x.get("region", ""))):
        cov.append([
            row.get("account_id", ""), row.get("account_name", ""), row.get("region", ""), row.get("service", ""),
            row.get("resource_type", ""), row.get("resource_count", 0), row.get("scan_status", ""), row.get("message", ""), row.get("source", ""),
        ])
    _style_header(cov)
    cov.freeze_panes = "A2"
    cov.auto_filter.ref = cov.dimensions
    for col, width in {"A": 16, "B": 28, "C": 18, "D": 22, "E": 34, "F": 15, "G": 18, "H": 36, "I": 24}.items():
        cov.column_dimensions[col].width = width

    # Preserve every API-returned attribute without forcing nested JSON into one Excel cell.
    # Values are flattened into path/value rows; oversized single values are split into parts.
    attr_sheet_index = 1
    attrs = wb.create_sheet("Resource Attributes 1")
    attrs.append(["Resource Key", "Account ID", "Region", "Service", "Resource Type", "Section", "Attribute Path", "Value"])
    _style_header(attrs)
    attr_row = 1
    for r in rows:
        sections = (
            ("Tags", r.tags), ("Configuration", r.configuration), ("Networking", r.networking),
            ("Security", r.security), ("Relationships", r.relationships), ("Raw API Attributes", r.raw_attributes),
        )
        for section_name, payload in sections:
            for path, value in _flatten(payload):
                if attr_row >= 1000000:
                    attrs.freeze_panes = "A2"
                    attr_sheet_index += 1
                    attrs = wb.create_sheet(f"Resource Attributes {attr_sheet_index}")
                    attrs.append(["Resource Key", "Account ID", "Region", "Service", "Resource Type", "Section", "Attribute Path", "Value"])
                    _style_header(attrs)
                    attr_row = 1
                attrs.append([r.resource_key, r.account_id, r.region, r.service, r.resource_type, section_name, path, value])
                attr_row += 1
    for sheet in [s for s in wb.worksheets if s.title.startswith("Resource Attributes")]:
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 16
        sheet.column_dimensions["C"].width = 16
        sheet.column_dimensions["D"].width = 20
        sheet.column_dimensions["E"].width = 30
        sheet.column_dimensions["F"].width = 22
        sheet.column_dimensions["G"].width = 60
        sheet.column_dimensions["H"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="c-sage-complete-aws-inventory.xlsx"'
    return response


def export_resource_xlsx(request, resource_key):
    scan = _latest()
    r = ResourceInventory.objects.by_key(scan, resource_key)
    if not r:
        return HttpResponse("Resource not found", status=404, content_type="text/plain")

    wb = Workbook()
    ws = wb.active
    ws.title = "Resource"
    ws.append(["Field", "Value"])
    _style_header(ws)
    for label, value in (
        ("Resource Key", r.resource_key), ("Account ID", r.account_id), ("Account Name", r.account_name),
        ("Region", r.region), ("Service", r.service), ("Resource Type", r.resource_type), ("Resource Name", r.resource_name),
        ("Resource ID", r.resource_id), ("ARN", r.resource_arn), ("Status", r.status),
        ("Creation Time", r.creation_time), ("Discovery Source", r.discovery_source),
    ):
        ws.append([label, str(value or "")])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100

    attrs = wb.create_sheet("Attributes")
    attrs.append(["Section", "Attribute Path", "Value"])
    _style_header(attrs)
    for section_name, payload in (
        ("Tags", r.tags), ("Configuration", r.configuration), ("Networking", r.networking),
        ("Security", r.security), ("Relationships", r.relationships), ("Raw API Attributes", r.raw_attributes),
    ):
        for path, value in _flatten(payload):
            attrs.append([section_name, path, value])
    attrs.freeze_panes = "A2"
    attrs.column_dimensions["A"].width = 24
    attrs.column_dimensions["B"].width = 70
    attrs.column_dimensions["C"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="c-sage-resource-{r.resource_key[:12]}.xlsx"'
    return response


def export_resource_json(request, resource_key):
    scan = _latest()
    r = ResourceInventory.objects.by_key(scan, resource_key)
    if not r:
        return JsonResponse({"error": "Resource not found"}, status=404)
    return JsonResponse(r.to_dict(), json_dumps_params={"indent": 2, "default": str})


def export_findings_csv(request):
    scan = _latest()
    rows = Finding.objects.for_scan(scan)
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="c-sage-compliance-findings.csv"'
    writer = csv.writer(response)
    writer.writerow(["Rule", "Module", "Account ID", "Account Name", "Region", "Service", "Resource ID", "Status", "Severity", "Title", "Details"])
    for f in rows:
        writer.writerow([f.rule_id, f.module, f.account_id, f.account_name, f.region, f.service, f.resource_id, f.status, f.severity, f.title, f.details])
    return response
