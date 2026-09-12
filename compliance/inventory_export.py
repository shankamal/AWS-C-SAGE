import io
import re
from collections import defaultdict
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ORANGE = "E36D1F"
NAVY = "003366"
WHITE = "FFFFFF"
BORDER = "D8E0E8"
EXCEL_MAX_COLUMNS = 16384
PROPERTY_COLUMNS_PER_SHEET = 15000
EXCEL_CELL_SAFE_LIMIT = 32000

BASE_COLUMNS = [
    ("Resource Key", "resource_key"),
    ("Account ID", "account_id"),
    ("Account Name", "account_name"),
    ("Region", "region"),
    ("Service", "service"),
    ("Resource Type", "resource_type"),
    ("Resource Name", "resource_name"),
    ("Resource ID", "resource_id"),
    ("ARN", "resource_arn"),
    ("Status", "status"),
    ("Creation Time", "creation_time"),
    ("Discovery Source", "discovery_source"),
]


def _style_header(ws, row=1):
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(color=WHITE, bold=True)
    border = Border(bottom=Side(style="thin", color=BORDER))
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[row].height = 26


def _safe_cell_parts(value):
    text = "" if value is None else str(value)
    if len(text) <= EXCEL_CELL_SAFE_LIMIT:
        return [text]
    return [text[i:i + EXCEL_CELL_SAFE_LIMIT] for i in range(0, len(text), EXCEL_CELL_SAFE_LIMIT)]


def _flatten_to_columns(value, prefix=""):
    """Flatten nested AWS API output into independent Excel columns."""
    result = {}
    if isinstance(value, dict):
        if not value and prefix:
            result[prefix] = "{}"
        for key, child in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            result.update(_flatten_to_columns(child, path))
    elif isinstance(value, list):
        if not value and prefix:
            result[prefix] = "[]"
        for index, child in enumerate(value):
            path = f"{prefix}[{index}]" if prefix else f"[{index}]"
            result.update(_flatten_to_columns(child, path))
    else:
        parts = _safe_cell_parts(value)
        if len(parts) == 1:
            result[prefix or "Value"] = parts[0]
        else:
            for index, part in enumerate(parts, start=1):
                result[f"{prefix or 'Value'} [part {index}]"] = part
    return result


def _resource_properties(resource):
    properties = {}
    for section_name, payload in (
        ("Tags", resource.tags),
        ("Configuration", resource.configuration),
        ("Networking", resource.networking),
        ("Security", resource.security),
        ("Relationships", resource.relationships),
        ("Raw API", resource.raw_attributes),
    ):
        properties.update(_flatten_to_columns(payload or {}, section_name))
    return properties


def _safe_sheet_name(value, used):
    clean = re.sub(r"[\\/*?:\[\]]", "-", value or "Resources").strip() or "Resources"
    clean = clean[:31]
    candidate = clean
    suffix = 2
    while candidate.lower() in used:
        tail = f"-{suffix}"
        candidate = f"{clean[:31-len(tail)]}{tail}"
        suffix += 1
    used.add(candidate.lower())
    return candidate


def _format_data_sheet(ws, property_count=0):
    _style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    base_widths = [22, 16, 26, 16, 20, 28, 30, 34, 48, 18, 24, 22]
    for index, width in enumerate(base_widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for index in range(len(BASE_COLUMNS) + 1, len(BASE_COLUMNS) + property_count + 1):
        ws.column_dimensions[get_column_letter(index)].width = 26
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_base_inventory_sheet(wb, rows, used_names):
    ws = wb.create_sheet(_safe_sheet_name("Filtered Inventory", used_names))
    ws.append([label for label, _ in BASE_COLUMNS])
    for resource in rows:
        ws.append([str(getattr(resource, attr, "") or "") for _, attr in BASE_COLUMNS])
    _format_data_sheet(ws, 0)


def _write_resource_type_sheets(wb, rows, used_names):
    grouped = defaultdict(list)
    for resource in rows:
        grouped[(resource.service or "Unknown", resource.resource_type or "Unknown")].append(resource)

    for (service, resource_type), resources in sorted(grouped.items(), key=lambda item: (item[0][0].lower(), item[0][1].lower())):
        property_maps = [_resource_properties(resource) for resource in resources]
        property_names = sorted({key for mapping in property_maps for key in mapping})
        chunks = [property_names[i:i + PROPERTY_COLUMNS_PER_SHEET] for i in range(0, len(property_names), PROPERTY_COLUMNS_PER_SHEET)] or [[]]

        for chunk_index, property_chunk in enumerate(chunks, start=1):
            label = f"{service}-{resource_type}"
            if len(chunks) > 1:
                label += f"-{chunk_index}"
            ws = wb.create_sheet(_safe_sheet_name(label, used_names))
            ws.append([label for label, _ in BASE_COLUMNS] + property_chunk)
            for resource, mapping in zip(resources, property_maps):
                base_values = [str(getattr(resource, attr, "") or "") for _, attr in BASE_COLUMNS]
                ws.append(base_values + [mapping.get(name, "") for name in property_chunk])
            _format_data_sheet(ws, len(property_chunk))


def _filter_coverage(coverage, filters):
    result = list(coverage)
    if filters.get("account"):
        result = [row for row in result if row.get("account_id") == filters["account"]]
    if filters.get("service"):
        result = [row for row in result if row.get("service") == filters["service"]]
    if filters.get("resource_type"):
        result = [row for row in result if row.get("resource_type") == filters["resource_type"]]
    if filters.get("region"):
        result = [row for row in result if row.get("region") == filters["region"]]
    return result


def build_inventory_workbook(scan, rows, coverage, filters):
    """Build a search-filtered workbook with one Excel column per AWS resource property."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Inventory Summary"
    summary["A1"] = "C-SAGE AWS Master Inventory"
    summary["A1"].font = Font(size=18, bold=True, color=WHITE)
    summary["A1"].fill = PatternFill("solid", fgColor=ORANGE)
    summary.merge_cells("A1:D1")

    applied = [f"{key}={value}" for key, value in filters.items() if value]
    filtered_coverage = _filter_coverage(coverage, filters)
    summary_rows = [
        ("Scan ID", scan.id if scan else ""),
        ("Generated At (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Resources Exported", len(rows)),
        ("Applied Search Criteria", "; ".join(applied) if applied else "All inventory resources"),
        ("Coverage Records", len(filtered_coverage)),
        ("Zero-resource Coverage Records", sum(1 for row in filtered_coverage if int(row.get("resource_count", 0) or 0) == 0)),
        ("Access/Error Coverage Records", sum(1 for row in filtered_coverage if row.get("scan_status") != "COMPLETE")),
        ("Export Model", "One row per resource; every nested AWS property flattened to its own column in service/resource-type sheets"),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=3):
        summary.cell(row=row_index, column=1, value=label)
        summary.cell(row=row_index, column=2, value=value)
        summary.cell(row=row_index, column=1).font = Font(bold=True, color=NAVY)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 100

    used_names = {"inventory summary"}
    _write_base_inventory_sheet(wb, rows, used_names)
    _write_resource_type_sheets(wb, rows, used_names)

    cov = wb.create_sheet(_safe_sheet_name("Service Coverage", used_names))
    cov.append(["Account ID", "Account Name", "Region", "Service", "Resource Type", "Resource Count", "Scan Status", "Message", "Source"])
    for row in sorted(filtered_coverage, key=lambda item: (item.get("account_name", ""), item.get("service", ""), item.get("resource_type", ""), item.get("region", ""))):
        cov.append([
            row.get("account_id", ""), row.get("account_name", ""), row.get("region", ""), row.get("service", ""),
            row.get("resource_type", ""), row.get("resource_count", 0), row.get("scan_status", ""), row.get("message", ""), row.get("source", ""),
        ])
    _style_header(cov)
    cov.freeze_panes = "A2"
    cov.auto_filter.ref = cov.dimensions
    for col, width in {"A": 16, "B": 28, "C": 18, "D": 22, "E": 34, "F": 15, "G": 18, "H": 44, "I": 24}.items():
        cov.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
